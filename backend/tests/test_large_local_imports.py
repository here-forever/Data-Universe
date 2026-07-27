from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.errors import AppError
from app.datasets.materializer import DatasetMaterializer
from app.imports.parser import ImportGuardrails, parse_tabular_path
from app.imports.schemas import ImportFieldPreview
from app.imports.storage import LocalFileStorage


def guardrails(**overrides: int | float) -> ImportGuardrails:
    values: dict[str, int | float] = {
        "max_file_size_bytes": 10_000_000,
        "max_rows": 10_000,
        "parse_timeout_seconds": 30,
        "inference_sample_size": 3,
        "preview_sample_size": 2,
    }
    values.update(overrides)
    return ImportGuardrails(**values)


def login(client: TestClient) -> dict[str, str]:
    response = client.post(
        "/api/auth/login",
        json={"email": "admin@example.com", "password": "admin123"},
    )
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def create_project(client: TestClient, headers: dict[str, str]) -> str:
    response = client.post(
        "/api/projects",
        headers=headers,
        json={"name": "Large Import Project", "description": None},
    )
    assert response.status_code == 201
    return response.json()["id"]


def test_csv_scan_keeps_bounded_sample_and_exact_row_count(tmp_path: Path) -> None:
    source_path = tmp_path / "large.csv"
    source_path.write_text(
        "order_id,amount\n" + "".join(f"{index},{index * 1.5}\n" for index in range(25)),
        encoding="utf-8",
    )

    parsed = parse_tabular_path("large.csv", source_path, guardrails())

    assert parsed.row_count == 25
    assert len(parsed.sample_rows) == 2
    assert [field.inferred_type for field in parsed.fields] == ["integer", "decimal"]
    assert len(list(parsed.iter_rows())) == 25
    assert len(list(parsed.iter_rows())) == 25


def test_excel_scan_is_reopenable_and_counts_all_rows(tmp_path: Path) -> None:
    source_path = tmp_path / "large.xlsx"
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet()
    worksheet.append(["order_id", "region"])
    for index in range(17):
        worksheet.append([index, "East" if index % 2 else "West"])
    workbook.save(source_path)

    parsed = parse_tabular_path("large.xlsx", source_path, guardrails())

    assert parsed.row_count == 17
    assert len(parsed.sample_rows) == 2
    assert sum(1 for _ in parsed.iter_rows()) == 17


@pytest.mark.parametrize(
    ("limits", "content", "expected_code"),
    [
        ({"max_file_size_bytes": 4}, b"value\n1\n", "file_size_limit_exceeded"),
        ({"max_rows": 1}, b"value\n1\n2\n", "file_row_limit_exceeded"),
    ],
)
def test_parser_guardrails_return_actionable_errors(
    tmp_path: Path,
    limits: dict[str, int],
    content: bytes,
    expected_code: str,
) -> None:
    source_path = tmp_path / "guarded.csv"
    source_path.write_bytes(content)

    with pytest.raises(AppError) as error:
        parse_tabular_path("guarded.csv", source_path, guardrails(**limits))

    assert error.value.code == expected_code
    assert "IMPORT_" in error.value.message


def test_parser_enforces_elapsed_time_guardrail(tmp_path: Path, monkeypatch) -> None:
    source_path = tmp_path / "slow.csv"
    source_path.write_text("value\n1\n", encoding="utf-8")
    timestamps = iter([0.0, 2.0])
    monkeypatch.setattr("app.imports.parser.monotonic", lambda: next(timestamps))

    with pytest.raises(AppError) as error:
        parse_tabular_path(
            "slow.csv",
            source_path,
            guardrails(parse_timeout_seconds=1),
        )

    assert error.value.code == "file_parse_timeout"
    assert "IMPORT_PARSE_TIMEOUT_SECONDS" in error.value.message


def test_upload_storage_copies_stream_in_bounded_chunks(tmp_path: Path) -> None:
    content = b"abcdefghijklmnopqrstuvwxyz"

    class TrackingStream(BytesIO):
        def __init__(self, value: bytes) -> None:
            super().__init__(value)
            self.requested_sizes: list[int] = []

        def read(self, size: int = -1) -> bytes:
            self.requested_sizes.append(size)
            return super().read(size)

    stream = TrackingStream(content)
    storage = LocalFileStorage(str(tmp_path / "uploads"), chunk_size_bytes=5)

    storage_path, size_bytes = storage.save_upload_stream(
        project_id="prj_test",
        uploaded_file_id="file_test",
        file_name="orders.csv",
        stream=stream,
    )

    assert size_bytes == len(content)
    assert Path(storage_path).read_bytes() == content
    assert stream.requested_sizes
    assert set(stream.requested_sizes) == {5}


def test_materializer_inserts_rows_in_configured_batches() -> None:
    engine = create_engine("sqlite+pysqlite:///:memory:")
    insert_batch_sizes: list[int] = []

    def record_insert_batch(
        _connection,
        _cursor,
        statement: str,
        parameters,
        _context,
        executemany: bool,
    ) -> None:
        if statement.lstrip().upper().startswith("INSERT INTO") and "batch_rows" in statement:
            insert_batch_sizes.append(len(parameters) if executemany else 1)

    event.listen(engine, "before_cursor_execute", record_insert_batch)
    fields = [ImportFieldPreview(name="value", inferred_type="integer", nullable=False, order=0)]
    checkpoints: list[int] = []

    with Session(engine) as session:
        inserted_count = DatasetMaterializer(session, batch_size=3).create_table(
            table_name="batch_rows",
            fields=fields,
            rows=({"value": index} for index in range(7)),
            progress_callback=checkpoints.append,
        )
        session.commit()

    assert inserted_count == 7
    assert insert_batch_sizes == [3, 3, 1]
    assert checkpoints == [3, 6, 7]


def test_size_guardrail_retains_staged_file_and_task_checkpoint(
    client: TestClient,
    monkeypatch,
) -> None:
    monkeypatch.setenv("IMPORT_MAX_FILE_SIZE_BYTES", "8")
    get_settings.cache_clear()
    headers = login(client)
    project_id = create_project(client, headers)
    content = b"value\n123456789\n"

    response = client.post(
        "/api/imports/file-previews",
        headers=headers,
        data={"project_id": project_id},
        files={"file": ("oversized.csv", content, "text/csv")},
    )

    assert response.status_code == 413
    assert response.json()["error"]["code"] == "file_size_limit_exceeded"
    uploads = client.get(
        "/api/imports/uploads",
        headers=headers,
        params={"project_id": project_id},
    ).json()["items"]
    assert uploads[0]["status"] == "failed"
    assert uploads[0]["size_bytes"] == len(content)
    staged_files = list(
        Path(get_settings().upload_storage_root).glob(f"{project_id}/*/oversized.csv")
    )
    assert len(staged_files) == 1
    assert staged_files[0].read_bytes() == content

    tasks = client.get(
        "/api/tasks",
        headers=headers,
        params={"project_id": project_id},
    ).json()["items"]
    parse_task = next(task for task in tasks if task["task_type"] == "file_preview_parse")
    assert parse_task["status"] == "failed"
    assert parse_task["progress"] == 25
    assert parse_task["can_retry"] is False


def test_materialization_failure_rolls_back_dataset_and_keeps_retry_context(
    client: TestClient,
    monkeypatch,
) -> None:
    monkeypatch.setenv("IMPORT_INFERENCE_SAMPLE_SIZE", "1")
    monkeypatch.setenv("IMPORT_MATERIALIZATION_BATCH_SIZE", "1")
    get_settings.cache_clear()
    headers = login(client)
    project_id = create_project(client, headers)
    upload_response = client.post(
        "/api/imports/file-previews",
        headers=headers,
        data={"project_id": project_id},
        files={"file": ("mixed.csv", b"value\n2\nnot-an-integer\n", "text/csv")},
    )
    assert upload_response.status_code == 201
    preview = upload_response.json()
    assert preview["fields"][0]["inferred_type"] == "integer"

    dataset_response = client.post(
        "/api/datasets",
        headers=headers,
        json={
            "project_id": project_id,
            "preview_id": preview["id"],
            "name": "Mixed Values",
            "fields": preview["fields"],
        },
    )

    assert dataset_response.status_code == 422
    assert dataset_response.json()["error"]["code"] == "file_row_type_mismatch"
    datasets = client.get(
        "/api/datasets",
        headers=headers,
        params={"project_id": project_id},
    ).json()["items"]
    assert datasets == []
    tasks = client.get(
        "/api/tasks",
        headers=headers,
        params={"project_id": project_id},
    ).json()["items"]
    materialization_task = next(
        task for task in tasks if task["task_type"] == "dataset_materialization"
    )
    assert materialization_task["status"] == "failed"
    assert materialization_task["progress"] == 35
    assert materialization_task["can_retry"] is True

    upload_root = Path(get_settings().upload_storage_root)
    assert next(upload_root.glob(f"{project_id}/*/mixed.csv")).exists()


def test_materialization_rejects_source_row_count_drift(
    client: TestClient,
) -> None:
    headers = login(client)
    project_id = create_project(client, headers)
    upload_response = client.post(
        "/api/imports/file-previews",
        headers=headers,
        data={"project_id": project_id},
        files={"file": ("orders.csv", b"order_id\n2\n3\n", "text/csv")},
    )
    assert upload_response.status_code == 201
    preview = upload_response.json()
    source_path = next(Path(get_settings().upload_storage_root).glob(f"{project_id}/*/orders.csv"))
    with source_path.open("a", encoding="utf-8") as source:
        source.write("4\n")

    dataset_response = client.post(
        "/api/datasets",
        headers=headers,
        json={
            "project_id": project_id,
            "preview_id": preview["id"],
            "name": "Changed Orders",
            "fields": preview["fields"],
        },
    )

    assert dataset_response.status_code == 409
    assert dataset_response.json()["error"]["code"] == "source_row_count_changed"
    datasets = client.get(
        "/api/datasets",
        headers=headers,
        params={"project_id": project_id},
    ).json()["items"]
    assert datasets == []
