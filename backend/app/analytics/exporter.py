import csv
import re
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.analytics.schemas import AnalysisResponse


def export_analysis_csv(result: AnalysisResponse) -> bytes:
    buffer = StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=result.columns, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(result.rows)
    return buffer.getvalue().encode("utf-8-sig")


def export_analysis_xlsx(result: AnalysisResponse) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Analysis"
    sheet.append(result.columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for row in result.rows:
        sheet.append([row.get(column) for column in result.columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        width = min(36, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))
        sheet.column_dimensions[column_cells[0].column_letter].width = width

    metadata = workbook.create_sheet("Metadata")
    metadata.append(["Dataset", result.dataset_name])
    metadata.append(["Source rows", result.source_row_count])
    metadata.append(["Filtered rows", result.filtered_row_count])
    metadata.append(["Groups", result.total_groups])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def safe_export_filename(dataset_name: str, suffix: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9_-]+", "_", dataset_name).strip("_")
    return f"{normalized or 'analysis'}_analysis.{suffix}"
