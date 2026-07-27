import csv
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO, StringIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.visualizations.schemas import (
    DashboardActiveSelection,
    DashboardGlobalFilter,
    ReportExportFormat,
)


@dataclass(frozen=True)
class ReportChartData:
    chart_id: str
    chart_name: str
    chart_type: str
    data_view_id: str
    dimension: str
    metric: str
    aggregation: str
    rows: list[dict[str, object | None]]


@dataclass(frozen=True)
class ReportDocument:
    export_id: str
    dashboard_id: str
    dashboard_name: str
    dashboard_mode: str
    dashboard_version: int
    generated_at: datetime
    filters: list[DashboardGlobalFilter]
    selections: list[DashboardActiveSelection]
    charts: list[ReportChartData]


@dataclass(frozen=True)
class AggregatedReportRow:
    label: str
    value: float


CONTENT_TYPES: dict[ReportExportFormat, str] = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def export_report(document: ReportDocument, export_format: ReportExportFormat) -> bytes:
    if export_format == "csv":
        return export_report_csv(document)
    if export_format == "xlsx":
        return export_report_xlsx(document)
    return export_report_pdf(document)


def export_report_csv(document: ReportDocument) -> bytes:
    buffer = StringIO(newline="")
    columns = [
        "export_id",
        "dashboard_id",
        "dashboard_name",
        "dashboard_version",
        "chart_id",
        "chart_name",
        "data_view_id",
        "chart_type",
        "dimension",
        "metric",
        "aggregation",
        "dimension_value",
        "metric_value",
    ]
    writer = csv.DictWriter(buffer, fieldnames=columns)
    writer.writeheader()
    for chart in document.charts:
        for row in aggregate_report_chart(chart, document.filters, document.selections):
            writer.writerow(
                {
                    "export_id": document.export_id,
                    "dashboard_id": document.dashboard_id,
                    "dashboard_name": document.dashboard_name,
                    "dashboard_version": document.dashboard_version,
                    "chart_id": chart.chart_id,
                    "chart_name": chart.chart_name,
                    "data_view_id": chart.data_view_id,
                    "chart_type": chart.chart_type,
                    "dimension": chart.dimension,
                    "metric": chart.metric,
                    "aggregation": chart.aggregation,
                    "dimension_value": row.label,
                    "metric_value": row.value,
                }
            )
    return buffer.getvalue().encode("utf-8-sig")


def export_report_xlsx(document: ReportDocument) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Report"
    summary_rows = [
        ("Export ID", document.export_id),
        ("Source dashboard ID", document.dashboard_id),
        ("Dashboard", document.dashboard_name),
        ("Mode", document.dashboard_mode),
        ("Configuration version", document.dashboard_version),
        ("Generated at", document.generated_at.isoformat()),
        ("Charts", len(document.charts)),
        ("Global filters", len(document.filters)),
        ("Active selections", len(document.selections)),
    ]
    for label, value in summary_rows:
        summary.append([label, value])
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 64
    for cell in summary["A"]:
        cell.font = Font(bold=True, color="334155")

    used_titles = {"Report"}
    for index, chart in enumerate(document.charts, start=1):
        title = unique_sheet_title(chart.chart_name, index, used_titles)
        used_titles.add(title)
        sheet = workbook.create_sheet(title)
        sheet.append(["Dimension", "Value"])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="6D5DFB")
            cell.alignment = Alignment(horizontal="center")
        for row in aggregate_report_chart(chart, document.filters, document.selections):
            sheet.append([row.label, row.value])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions["A"].width = 34
        sheet.column_dimensions["B"].width = 18

        metadata = workbook.create_sheet(unique_sheet_title(f"{index} Source", index, used_titles))
        used_titles.add(metadata.title)
        metadata.append(["Chart ID", chart.chart_id])
        metadata.append(["Data view ID", chart.data_view_id])
        metadata.append(["Chart type", chart.chart_type])
        metadata.append(["Dimension", chart.dimension])
        metadata.append(["Metric", chart.metric])
        metadata.append(["Aggregation", chart.aggregation])
        metadata.column_dimensions["A"].width = 20
        metadata.column_dimensions["B"].width = 48

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_report_pdf(document: ReportDocument) -> bytes:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    buffer = BytesIO()
    page_size = landscape(A4) if document.dashboard_mode == "screen" else A4
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=document.dashboard_name,
        author="Data Analysis System",
        subject=f"Dashboard export {document.dashboard_id}",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=30,
        textColor=colors.HexColor("#2F2A55"),
        alignment=TA_CENTER,
        spaceAfter=8 * mm,
    )
    heading_style = ParagraphStyle(
        "ReportHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=21,
        textColor=colors.HexColor("#3F3A70"),
        spaceAfter=4 * mm,
    )
    body_style = ParagraphStyle(
        "ReportBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor("#475569"),
    )
    cjk_title_style = ParagraphStyle(
        "ReportTitleCjk",
        parent=title_style,
        fontName="STSong-Light",
    )
    cjk_heading_style = ParagraphStyle(
        "ReportHeadingCjk",
        parent=heading_style,
        fontName="STSong-Light",
    )
    cjk_body_style = ParagraphStyle(
        "ReportBodyCjk",
        parent=body_style,
        fontName="STSong-Light",
    )
    header_style = ParagraphStyle(
        "ReportTableHeader",
        parent=body_style,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )
    cjk_header_style = ParagraphStyle(
        "ReportTableHeaderCjk",
        parent=header_style,
        fontName="STSong-Light",
    )

    story: list[object] = [
        styled_paragraph(document.dashboard_name, title_style, cjk_title_style),
        styled_paragraph(
            f"Printable {document.dashboard_mode} export - source {document.dashboard_id} "
            f"- version {document.dashboard_version}",
            body_style,
            cjk_body_style,
        ),
        Spacer(1, 5 * mm),
        metadata_table(document, body_style, cjk_body_style),
        Spacer(1, 8 * mm),
    ]
    for index, chart in enumerate(document.charts):
        if index > 0:
            story.append(PageBreak())
        story.append(styled_paragraph(chart.chart_name, heading_style, cjk_heading_style))
        story.append(
            styled_paragraph(
                f"Chart {chart.chart_id} - data view {chart.data_view_id} - "
                f"{chart.aggregation.upper()}({chart.metric}) by {chart.dimension}",
                body_style,
                cjk_body_style,
            )
        )
        story.append(Spacer(1, 4 * mm))
        aggregated = aggregate_report_chart(chart, document.filters, document.selections)
        story.append(
            chart_data_table(
                aggregated,
                body_style,
                cjk_body_style,
                header_style,
                cjk_header_style,
            )
        )

    pdf.build(story, onFirstPage=draw_page_footer, onLaterPages=draw_page_footer)
    return buffer.getvalue()


def aggregate_report_chart(
    chart: ReportChartData,
    filters: list[DashboardGlobalFilter],
    selections: list[DashboardActiveSelection],
) -> list[AggregatedReportRow]:
    groups: dict[str, tuple[float, int]] = {}
    for row in chart.rows:
        if not row_matches_filters(row, chart, filters, selections):
            continue
        label = format_value(row.get(chart.dimension))
        current_sum, current_count = groups.get(label, (0.0, 0))
        numeric_value = 1.0 if chart.aggregation == "count" else to_number(row.get(chart.metric))
        groups[label] = (current_sum + numeric_value, current_count + 1)

    result: list[AggregatedReportRow] = []
    for label, (total, count) in groups.items():
        value = total / count if chart.aggregation == "avg" and count else total
        result.append(AggregatedReportRow(label=label, value=round(value, 2)))
    return result


def row_matches_filters(
    row: dict[str, object | None],
    chart: ReportChartData,
    filters: list[DashboardGlobalFilter],
    selections: list[DashboardActiveSelection],
) -> bool:
    for item in filters:
        if item.value is None or item.value == "":
            continue
        if item.data_view_id and item.data_view_id != chart.data_view_id:
            continue
        if item.field not in row or not compare_values(
            row.get(item.field), item.operator, item.value
        ):
            return False
    for selection in selections:
        if selection.chart_id == chart.chart_id or selection.field not in row:
            continue
        if not compare_values(row.get(selection.field), "eq", selection.value):
            return False
    return True


def compare_values(left: object | None, operator: str, right: object | None) -> bool:
    if operator == "contains":
        return str(right or "").casefold() in str(left or "").casefold()
    if operator in {"gt", "gte", "lt", "lte"}:
        left_number = to_number(left)
        right_number = to_number(right)
        if operator == "gt":
            return left_number > right_number
        if operator == "gte":
            return left_number >= right_number
        if operator == "lt":
            return left_number < right_number
        return left_number <= right_number
    equal = normalize_comparable(left) == normalize_comparable(right)
    return not equal if operator == "neq" else equal


def safe_report_filename(name: str, dashboard_id: str, suffix: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9_-]+", "_", name).strip("_")
    stable_id = re.sub(r"[^A-Za-z0-9_-]+", "_", dashboard_id).strip("_")
    return f"{normalized or 'report'}_{stable_id}.{suffix}"


def unique_sheet_title(name: str, index: int, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", " ", name).strip() or f"Chart {index}"
    base = base[:31]
    candidate = base
    counter = 2
    while candidate in used:
        suffix = f" {counter}"
        candidate = f"{base[: 31 - len(suffix)]}{suffix}"
        counter += 1
    return candidate


def metadata_table(
    document: ReportDocument,
    body_style: ParagraphStyle,
    cjk_body_style: ParagraphStyle,
) -> Table:
    rows = [
        ["Export ID", document.export_id],
        ["Source dashboard", document.dashboard_id],
        ["Generated", document.generated_at.astimezone(UTC).isoformat()],
        ["Charts", str(len(document.charts))],
        ["Filters / selections", f"{len(document.filters)} / {len(document.selections)}"],
    ]
    table = Table(
        [[styled_paragraph(str(cell), body_style, cjk_body_style) for cell in row] for row in rows],
        colWidths=[42 * mm, 120 * mm],
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EDE9FE")),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#4C3F91")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D8D3EA")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def chart_data_table(
    rows: list[AggregatedReportRow],
    body_style: ParagraphStyle,
    cjk_body_style: ParagraphStyle,
    header_style: ParagraphStyle,
    cjk_header_style: ParagraphStyle,
) -> Table:
    data = [
        [
            styled_paragraph("Dimension", header_style, cjk_header_style),
            styled_paragraph("Value", header_style, cjk_header_style),
        ]
    ]
    for row in rows[:100]:
        data.append(
            [
                styled_paragraph(row.label, body_style, cjk_body_style),
                styled_paragraph(format(row.value, ",.2f"), body_style, cjk_body_style),
            ]
        )
    if len(data) == 1:
        data.append(
            [
                styled_paragraph("No matching data", body_style, cjk_body_style),
                styled_paragraph("-", body_style, cjk_body_style),
            ]
        )
    table = Table(data, colWidths=[115 * mm, 45 * mm], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6D5DFB")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D8D3EA")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def draw_page_footer(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawRightString(document.pagesize[0] - 18 * mm, 10 * mm, f"Page {document.page}")
    canvas.restoreState()


def styled_paragraph(
    value: str,
    latin_style: ParagraphStyle,
    cjk_style: ParagraphStyle,
) -> Paragraph:
    style = cjk_style if re.search(r"[\u3400-\u9fff]", value) else latin_style
    return Paragraph(escape(value), style)


def normalize_comparable(value: object | None) -> str:
    if value is None:
        return ""
    return str(value).casefold()


def format_value(value: object | None) -> str:
    if value is None or value == "":
        return "NULL"
    return str(value)


def to_number(value: object | None) -> float:
    if isinstance(value, bool) or value is None:
        return 0.0
    if isinstance(value, int | float):
        return float(value)
    try:
        return float(str(value))
    except ValueError:
        return 0.0
