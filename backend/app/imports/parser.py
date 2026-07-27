import csv
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, TextIOWrapper
from pathlib import Path
from time import monotonic

from openpyxl import load_workbook

from app.core.errors import AppError
from app.imports.schemas import FieldType, ImportFieldPreview

SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xlsm"}


@dataclass(frozen=True)
class ImportGuardrails:
    max_file_size_bytes: int
    max_rows: int
    parse_timeout_seconds: float
    inference_sample_size: int
    preview_sample_size: int


class ParsedTabularFile:
    def __init__(
        self,
        *,
        file_type: str,
        fields: list[ImportFieldPreview],
        row_count: int,
        sample_rows: list[dict[str, object | None]],
        row_iterator_factory: Callable[[], Iterator[dict[str, object | None]]],
    ) -> None:
        self.file_type = file_type
        self.fields = fields
        self.row_count = row_count
        self.sample_rows = sample_rows
        self._row_iterator_factory = row_iterator_factory

    def iter_rows(self) -> Iterator[dict[str, object | None]]:
        return self._row_iterator_factory()


def parse_tabular_file(
    file_name: str,
    content: bytes,
    guardrails: ImportGuardrails,
) -> ParsedTabularFile:
    validate_file_size(len(content), guardrails)
    return _analyze_tabular_source(
        file_name=file_name,
        source=content,
        guardrails=guardrails,
    )


def parse_tabular_path(
    file_name: str,
    path: Path,
    guardrails: ImportGuardrails,
) -> ParsedTabularFile:
    validate_file_size(path.stat().st_size, guardrails)
    return _analyze_tabular_source(
        file_name=file_name,
        source=path,
        guardrails=guardrails,
    )


def open_parsed_tabular_path(
    *,
    file_name: str,
    path: Path,
    fields: list[ImportFieldPreview],
    row_count: int,
    sample_rows: list[dict[str, object | None]],
    guardrails: ImportGuardrails,
) -> ParsedTabularFile:
    extension = validate_extension(file_name)
    validate_file_size(path.stat().st_size, guardrails)
    headers = [field.name for field in sorted(fields, key=lambda field: field.order)]
    return ParsedTabularFile(
        file_type=extension.removeprefix("."),
        fields=fields,
        row_count=row_count,
        sample_rows=sample_rows,
        row_iterator_factory=lambda: iter_typed_rows(
            file_name=file_name,
            source=path,
            headers=headers,
            fields=fields,
            guardrails=guardrails,
        ),
    )


def _analyze_tabular_source(
    *,
    file_name: str,
    source: Path | bytes,
    guardrails: ImportGuardrails,
) -> ParsedTabularFile:
    extension = validate_extension(file_name)
    started_at = monotonic()
    raw_rows = iter_raw_rows(file_name=file_name, source=source)
    try:
        raw_headers = next(raw_rows)
    except StopIteration as error:
        raise AppError(
            message="Uploaded file has no rows",
            code="empty_file",
            status_code=400,
        ) from error

    headers = normalize_headers(raw_headers)
    inference_rows: list[dict[str, object | None]] = []
    row_count = 0
    for raw_row in raw_rows:
        check_parse_timeout(started_at, guardrails)
        row = coerce_row(headers, raw_row)
        if not any(value is not None for value in row.values()):
            continue
        row_count += 1
        if row_count > guardrails.max_rows:
            raise AppError(
                message=(
                    f"File contains more than the configured {guardrails.max_rows:,} row limit. "
                    "Reduce the file or increase IMPORT_MAX_ROWS."
                ),
                code="file_row_limit_exceeded",
                status_code=413,
            )
        if len(inference_rows) < guardrails.inference_sample_size:
            inference_rows.append(row)

    check_parse_timeout(started_at, guardrails)
    fields = infer_fields(headers, inference_rows)
    typed_sample_rows = [
        coerce_typed_row(row, fields) for row in inference_rows[: guardrails.preview_sample_size]
    ]
    return ParsedTabularFile(
        file_type=extension.removeprefix("."),
        fields=fields,
        row_count=row_count,
        sample_rows=typed_sample_rows,
        row_iterator_factory=lambda: iter_typed_rows(
            file_name=file_name,
            source=source,
            headers=headers,
            fields=fields,
            guardrails=guardrails,
        ),
    )


def iter_typed_rows(
    *,
    file_name: str,
    source: Path | bytes,
    headers: list[str],
    fields: list[ImportFieldPreview],
    guardrails: ImportGuardrails,
) -> Iterator[dict[str, object | None]]:
    started_at = monotonic()
    raw_rows = iter_raw_rows(file_name=file_name, source=source)
    next(raw_rows, None)
    row_count = 0
    for raw_row in raw_rows:
        check_parse_timeout(started_at, guardrails)
        row = coerce_row(headers, raw_row)
        if not any(value is not None for value in row.values()):
            continue
        row_count += 1
        if row_count > guardrails.max_rows:
            raise AppError(
                message=(
                    f"File contains more than the configured {guardrails.max_rows:,} row limit. "
                    "Reduce the file or increase IMPORT_MAX_ROWS."
                ),
                code="file_row_limit_exceeded",
                status_code=413,
            )
        try:
            yield coerce_typed_row(row, fields)
        except (InvalidOperation, TypeError, ValueError) as error:
            raise AppError(
                message=(
                    f"Row {row_count:,} does not match the inferred field types. "
                    "Adjust the dataset field types or correct the source values and retry."
                ),
                code="file_row_type_mismatch",
                status_code=422,
            ) from error
    check_parse_timeout(started_at, guardrails)


def iter_raw_rows(
    *,
    file_name: str,
    source: Path | bytes,
) -> Iterator[list[object | None]]:
    extension = validate_extension(file_name)
    if extension == ".csv":
        yield from iter_csv_rows(source)
        return
    yield from iter_excel_rows(source)


def iter_csv_rows(source: Path | bytes) -> Iterator[list[object | None]]:
    if isinstance(source, Path):
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            for row in csv.reader(handle):
                yield list(row)
        return

    with TextIOWrapper(BytesIO(source), encoding="utf-8-sig", newline="") as handle:
        for row in csv.reader(handle):
            yield list(row)


def iter_excel_rows(source: Path | bytes) -> Iterator[list[object | None]]:
    workbook = load_workbook(
        filename=source if isinstance(source, Path) else BytesIO(source),
        read_only=True,
        data_only=True,
    )
    try:
        sheet = workbook.active
        for row in sheet.iter_rows(values_only=True):
            yield list(row)
    finally:
        workbook.close()


def validate_extension(file_name: str) -> str:
    extension = Path(file_name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise AppError(
            message="Only CSV and Excel files are supported",
            code="unsupported_file_type",
            status_code=400,
        )
    return extension


def validate_file_size(size_bytes: int, guardrails: ImportGuardrails) -> None:
    if size_bytes <= guardrails.max_file_size_bytes:
        return
    raise AppError(
        message=(
            f"File size exceeds the configured {guardrails.max_file_size_bytes:,} byte limit. "
            "Reduce the file or increase IMPORT_MAX_FILE_SIZE_BYTES."
        ),
        code="file_size_limit_exceeded",
        status_code=413,
    )


def check_parse_timeout(started_at: float, guardrails: ImportGuardrails) -> None:
    if monotonic() - started_at <= guardrails.parse_timeout_seconds:
        return
    raise AppError(
        message=(
            f"File parsing exceeded the configured {guardrails.parse_timeout_seconds:g} second "
            "limit. Reduce the file or increase IMPORT_PARSE_TIMEOUT_SECONDS."
        ),
        code="file_parse_timeout",
        status_code=408,
    )


def normalize_headers(raw_headers: list[object | None]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}

    for index, raw_header in enumerate(raw_headers):
        fallback = f"column_{index + 1}"
        name = str(raw_header).strip() if raw_header is not None else fallback
        name = name or fallback
        count = seen.get(name, 0)
        seen[name] = count + 1
        headers.append(name if count == 0 else f"{name}_{count + 1}")

    return headers


def coerce_row(headers: list[str], raw_row: list[object | None]) -> dict[str, object | None]:
    return {
        header: normalize_cell(raw_row[index]) if index < len(raw_row) else None
        for index, header in enumerate(headers)
    }


def normalize_cell(value: object | None) -> object | None:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value


def infer_fields(
    headers: list[str],
    rows: list[dict[str, object | None]],
) -> list[ImportFieldPreview]:
    fields: list[ImportFieldPreview] = []
    for order, header in enumerate(headers):
        values = [row[header] for row in rows]
        non_empty_values = [value for value in values if value is not None]
        fields.append(
            ImportFieldPreview(
                name=header,
                inferred_type=infer_type(non_empty_values),
                nullable=len(non_empty_values) != len(values),
                order=order,
            )
        )
    return fields


def infer_type(values: list[object]) -> FieldType:
    if not values:
        return "text"
    if all(is_boolean(value) for value in values):
        return "boolean"
    if all(is_integer(value) for value in values):
        return "integer"
    if all(is_decimal(value) for value in values):
        return "decimal"
    if all(is_date(value) for value in values):
        return "date"
    if all(is_datetime(value) for value in values):
        return "datetime"
    return "text"


def is_boolean(value: object) -> bool:
    if isinstance(value, bool):
        return True
    return isinstance(value, str) and value.lower() in {"true", "false", "yes", "no", "1", "0"}


def is_integer(value: object) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, int):
        return True
    return isinstance(value, str) and value.isdecimal()


def is_decimal(value: object) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, int | float | Decimal):
        return True
    if not isinstance(value, str):
        return False
    try:
        Decimal(value)
    except InvalidOperation:
        return False
    return True


def is_date(value: object) -> bool:
    if isinstance(value, datetime):
        return False
    if isinstance(value, date):
        return True
    if not isinstance(value, str):
        return False
    try:
        datetime.strptime(value, "%Y-%m-%d")
    except ValueError:
        return False
    return True


def is_datetime(value: object) -> bool:
    if isinstance(value, datetime):
        return True
    if not isinstance(value, str):
        return False
    try:
        datetime.fromisoformat(value)
    except ValueError:
        return False
    return True


def coerce_typed_row(
    row: dict[str, object | None],
    fields: list[ImportFieldPreview],
) -> dict[str, object | None]:
    return {field.name: coerce_value(row[field.name], field.inferred_type) for field in fields}


def coerce_value(value: object | None, field_type: FieldType) -> object | None:
    if value is None:
        return None
    if field_type == "integer":
        return int(value)
    if field_type == "decimal":
        return float(value)
    if field_type == "boolean":
        return coerce_boolean(value)
    if field_type == "date":
        return value.isoformat() if isinstance(value, date) else str(value)
    if field_type == "datetime":
        return value.isoformat() if isinstance(value, datetime) else str(value)
    return str(value)


def coerce_boolean(value: object) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).lower() in {"true", "yes", "1"}
